"""Parse hackathon store layout Excel (read-only from dataset/)."""

from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

ZONE_NAME_PATTERN = re.compile(
    r"(SKINCARE|MAKEUP|FRAGRANCE|HAIR|BATH|PERSONAL|SKIN|FRAGRANCE|BILLING|ENTRY)",
    re.I,
)


def find_layout_xlsx(data_dir: Path) -> Path | None:
    for pattern in ("*layout*.xlsx", "*Layout*.xlsx", "*.xlsx"):
        matches = sorted(data_dir.glob(pattern))
        if matches:
            return matches[0]
    return None


def _default_polygons() -> dict[str, list[list[int]]]:
    """1080p placeholders until calibrated per camera."""
    return {
        "SKINCARE": [[100, 150], [500, 150], [500, 400], [100, 400]],
        "MAKEUP": [[520, 150], [900, 150], [900, 400], [520, 400]],
        "FRAGRANCE": [[100, 420], [900, 420], [900, 650], [100, 650]],
    }


def parse_layout_xlsx(path: Path, store_id: str = "ST1008") -> dict[str, Any]:
    """
    Extract zone names from Excel; merge with default polygons in config.
    Hackathon xlsx may be descriptive (zone names per camera) not pixel polygons.
    """
    try:
        import openpyxl
    except ImportError as e:
        raise ImportError("pip install openpyxl to read store layout xlsx") from e

    zone_names: set[str] = set()
    open_hours: str | None = None
    camera_notes: dict[str, list[str]] = {}

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None]
            text = " ".join(cells).upper()
            for match in ZONE_NAME_PATTERN.finditer(text):
                name = match.group(1).upper()
                if name == "SKIN":
                    name = "SKINCARE"
                zone_names.add(name)
            if any("hour" in c.lower() for c in cells if len(c) < 30):
                open_hours = text[:200]

    wb.close()

    zones = _default_polygons()
    for z in zone_names:
        if z not in zones and z in ("SKINCARE", "MAKEUP", "FRAGRANCE", "HAIR", "BATH"):
            zones.setdefault(z, [[200, 200], [400, 200], [400, 400], [200, 400]])

    layout = {
        "store_id": store_id,
        "store_name": "Brigade_Bangalore",
        "zones": zones,
        "entry_line": {"start": [80, 540], "end": [400, 540], "in_direction": "down"},
        "billing_queue_zone": [[700, 300], [1050, 300], [1050, 700], [700, 700]],
        "billing_counter_zone": [[750, 150], [1050, 150], [1050, 300], [750, 300]],
        "source_xlsx": path.name,
        "open_hours_note": open_hours,
    }
    logger.info("Parsed layout from %s — zones: %s", path.name, list(zones.keys()))
    return layout


def cache_layout_json(layout: dict[str, Any], cache_path: Path) -> None:
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    with cache_path.open("w", encoding="utf-8") as f:
        json.dump(layout, f, indent=2)
